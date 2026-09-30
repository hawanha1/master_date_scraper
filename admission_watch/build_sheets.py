#!/usr/bin/env python3
"""
One-time builder for italy_universities.xlsx and germany_universities.xlsx.

- names come from ../Schengen_29_Public_Universities_complete.xlsx (duplicates removed)
- official domains are reused from ../intake_dates_<Country>.xlsx when available, otherwise looked
  up (Hipolabs universities API, then a web search)
- the admissions page is found by opening the homepage and picking the best admission link
- an existing sheet's Status / Notes / hand-edited links are preserved when rebuilding

Usage:  .venv/bin/python build_sheets.py [--country Italy]
"""
from __future__ import annotations

import argparse
import concurrent.futures as cf
import re
import sys
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT))
from registry import load_registry  # noqa: E402  (parent project)

from admission_watch import NOT_ADMISSION_PAGE, base_domain, page_text  # noqa: E402
from fetcher import Fetcher  # noqa: E402

COUNTRIES = {"Italy": "italy_universities.xlsx", "Germany": "germany_universities.xlsx"}
HEADER = ["#", "Country", "University", "Official Website", "Admissions Page", "Status", "Notes"]
ADMISSION_LINK = {
    "Italy": re.compile(r"immatricol|ammission|iscriver|iscrizion|futuri studenti|future students|admission"
                        r"|apply|how to apply|corsi di laurea|offerta formativa|studiare|international", re.I),
    "Germany": re.compile(r"bewerb|zulassung|einschreib|studieninteress|studienangebot|studium|admission"
                          r"|apply|application|prospective|international", re.I),
}
STRONG = re.compile(r"immatricol|ammission|bewerb|zulassung|admission|apply|application", re.I)
SKIP_DOMAINS = ("wikipedia", "facebook", "linkedin", "instagram", "youtube", "topuniversities", "shiksha",
                "yocket", "mastersportal", "studyportals", "unirank", "4icu", "timeshighereducation", "daad.de",
                "universitaly", "study-in-germany", "hochschulkompass", "leverageedu", "collegedunia")


# checked by hand: the earlier search-based lookup picked the wrong site for these
OVERRIDES = {
    "University of Chieti-Pescara": "unich.it",
    "University for Foreigners of Siena": "unistrasi.it",
    "Polytechnic University of Turin": "polito.it",
    "Gran Sasso Science Institute": "gssi.it",
    "University of Konstanz": "uni-konstanz.de",
    "Brandenburg University of Technology Cottbus-Senftenberg": "b-tu.de",
    "Ludwig Maximilian University of Munich": "lmu.de",
}


def old_domains(country: str) -> dict[str, str]:
    path = ROOT / f"intake_dates_{country}.xlsx"
    if not path.exists():
        return {}
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["Summary"]
    rows = list(ws.iter_rows(values_only=True))
    h = [str(x) for x in rows[0]]
    iu, iw = h.index("University"), h.index("Official Website")
    return {r[iu]: r[iw] for r in rows[1:] if r[iu] and r[iw]}


def lookup_domain(names: list[str], country: str) -> str | None:
    for name in names:
        try:
            data = requests.get("http://universities.hipolabs.com/search",
                                params={"name": name, "country": country}, timeout=15).json()
            exact = [d for d in data if d["name"].lower() == name.lower() and d.get("domains")]
            if exact:
                return exact[0]["domains"][0]
        except Exception:  # noqa: BLE001
            pass
    try:
        from ddgs import DDGS
        for r in DDGS().text(f"{names[0]} {country} official website", max_results=8):
            host = urlparse(r.get("href", "")).netloc.lower()
            if host and not any(s in host for s in SKIP_DOMAINS):
                return host
    except Exception:  # noqa: BLE001
        pass
    return None


def homepage_and_admissions(domain: str, country: str, fetcher: Fetcher) -> tuple[str, str, str]:
    root = base_domain(domain)
    for home in (f"https://www.{root}/", f"https://{root}/", f"https://{domain}/"):
        page = fetcher.get(home)
        if page.ok:
            break
    else:
        return f"https://www.{root}/", "", f"homepage: {page.status}"
    home = page.final_url or home
    _, links = page_text(page)
    rx = ADMISSION_LINK[country]
    best, best_score = "", 0
    for url, text in links:
        p = urlparse(url)
        if base_domain(p.netloc) != root or url.rstrip("/") == home.rstrip("/"):
            continue
        blob = f"{text} {p.path}"
        if not rx.search(blob) or NOT_ADMISSION_PAGE.search(url):
            continue
        score = 2 + 2 * bool(STRONG.search(blob)) + bool(re.search(r"/en/|english", url, re.I))
        score -= p.path.count("/") > 5
        if score > best_score:
            best, best_score = url.split("#")[0], score
    return home, best, "" if best else "admissions link not found on homepage, homepage is used"


def build(country: str):
    reg = load_registry(ROOT / "Schengen_29_Public_Universities_complete.xlsx")
    entries = [e for e in reg[country] if not e.duplicate_of]
    known = old_domains(country)
    out = HERE / COUNTRIES[country]
    keep = {}
    if out.exists():   # keep what the user typed
        ws = openpyxl.load_workbook(out).active
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[2]:
                keep[r[2]] = {"site": r[3], "adm": r[4], "status": r[5], "notes": r[6]}
    fetcher = Fetcher(country)

    def one(e):
        prev = keep.get(e.name, {})
        override = OVERRIDES.get(e.name)
        if override and override not in str(prev.get("site") or ""):
            prev = {k: v for k, v in prev.items() if k in ("status", "notes")}
        if prev.get("site") and prev.get("adm"):
            return e, prev["site"], prev["adm"], ""
        dom = override or known.get(e.name) or lookup_domain(e.aliases, country)
        if not dom:
            return e, "", "", "official website not found, please fill in"
        home, adm, note = homepage_and_admissions(dom.removeprefix("www."), country, fetcher)
        return e, prev.get("site") or home, prev.get("adm") or adm, note

    with cf.ThreadPoolExecutor(12) as ex:
        rows = list(ex.map(one, entries))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = country
    ws.append(HEADER)
    for i, (e, site, adm, note) in enumerate(rows, 1):
        prev = keep.get(e.name, {})
        ws.append([i, country, e.name, site, adm, prev.get("status") or "", prev.get("notes") or note])
        print(f"{i:>3}. {e.name[:45]:45} {site:35} {adm[:70]}")
    fill = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), fill, Alignment(vertical="center")
    for col, w in zip("ABCDEFG", (5, 10, 45, 35, 60, 12, 45)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"Applied,Skip"', allow_blank=True,
                        promptTitle="Status", prompt="Applied or Skip = not checked any more. Empty = checked nightly.")
    ws.add_data_validation(dv)
    dv.add(f"F2:F{len(rows) + 1}")
    wb.save(out)
    found = sum(1 for r in rows if r[2])
    print(f"\n{out.name}: {len(rows)} universities, admissions page found for {found}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--country", choices=list(COUNTRIES))
    a = ap.parse_args()
    for c in ([a.country] if a.country else COUNTRIES):
        build(c)
