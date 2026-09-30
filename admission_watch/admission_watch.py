#!/usr/bin/env python3
"""
Nightly admission watcher.

For every university in the country sheets (italy_universities.xlsx, germany_universities.xlsx)
whose Status is not "Applied" (or "Skip"), it opens the university's official admissions page,
follows admission-related links on the same website (bando / scadenze / Bewerbung / Fristen /
PDF calls ...), and looks for an application window that contains today's date.

Then it:
  1. writes open_admissions_<date>.xlsx   (universities whose admission is open now)
  2. writes scrape_problems_<date>.xlsx   (universities that could not be checked)
  3. emails them with Gmail (SMTP + app password)
  4. deletes both files, so nothing is left on the server

Run by cron at midnight (see install_cron.sh). Manual run:  .venv/bin/python admission_watch.py
Options: --dry-run (no email, prints results)  --keep (keep the sheets in ./reports)
         --only "Bologna,RWTH" (just these)    --country Italy
         --verify (open a visible browser so you can pass sites' human-verification checks once;
                   the nightly run reuses that browser session while the site accepts it)
"""
from __future__ import annotations

import argparse
import concurrent.futures as cf
import io
import logging
import os
import random
import re
import shutil
import smtplib
import ssl
import sys
import tempfile
import time
from dataclasses import dataclass, field
from datetime import date, datetime
from email.message import EmailMessage
from pathlib import Path
from urllib.parse import urljoin, urlparse

import openpyxl
from bs4 import BeautifulSoup
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader

import browser
from fetcher import Fetcher, Page, RobotsPolicy, search, wayback
from windows import Window, find_windows

HERE = Path(__file__).resolve().parent
log = logging.getLogger("watch")
logging.getLogger("pypdf").setLevel(logging.ERROR)

SKIP_STATUSES = {"applied", "skip", "ignore", "done"}
LINK_STRONG = re.compile(
    r"scadenz|bando|bandi|termin|frist|deadline|dates|calendar|calendario|immatricolazion|ammission"
    r"|preiscrizion|bewerbungsfrist|bewerbungszeitraum|zulassung|admission|application|how-to-apply|apply", re.I)
LINK_MEDIUM = re.compile(
    r"iscrizion|iscriversi|bewerb|einschreib|immatrikul|studieninteress|international|master|laurea|"
    r"degree|programm|studiengang|studium|study|studiare|future-students|futuri-studenti|prospective", re.I)
LINK_BAD = re.compile(
    r"login|logout|signin|password|cookie|privacy|impressum|imprint|datenschutz|facebook|instagram|twitter|"
    r"linkedin|youtube|tiktok|mailto:|tel:|javascript:|\.(?:jpg|jpeg|png|gif|svg|zip|docx?|xlsx?|pptx?|mp4)$"
    r"|/news/?$|jobs?|stellen|lavora|concorsi|personale|staff|research|ricerca|press", re.I)
# pages that carry dates but are not about admission (events, news, competitions, exchange, graduation)
NOT_ADMISSION_PAGE = re.compile(
    r"veranstaltung|/events?[/-]|/eventi|news-and-events|meldungen|aktuelles|/news/|/notizie/(?!immatricol)"
    r"|open-calls|funding|foerder|förder"
    r"|ausschreibung|wettbewerb|erasmus|mobilit|/internazionale/bandi|esame-di-laurea|allesame-di-laurea"
    r"|/laurearsi|dottorato|/phd|doctoral|promotion|stipendi|borse|/formulare|kontaktdaten|/career|/karriere", re.I)


# --------------------------------------------------------------------------- config
def load_config() -> dict:
    cfg = {"DEADLINE_DAYS": "30", "OPENED_DAYS": "45", "WORKERS": "12", "MAX_PAGES": "12",
           "SHEETS": "italy_universities.xlsx,germany_universities.xlsx", "SEND_WHEN_EMPTY": "false", "RESPECT_ROBOTS": "true"}
    path = HERE / "config.env"
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                cfg[k.strip()] = v.strip().strip('"').strip("'")
    for k in list(cfg) + ["GMAIL_ADDRESS", "GMAIL_APP_PASSWORD", "NOTIFY_TO", "PROXIES"]:
        if os.environ.get(k):
            cfg[k] = os.environ[k]
    return cfg


# --------------------------------------------------------------------------- sheets
@dataclass
class Uni:
    country: str
    name: str
    website: str
    admission_pages: list[str]
    status: str
    row: int


def read_sheet(path: Path) -> list[Uni]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(h or "").strip().lower() for h in rows[0]]

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    ci, cn, cw, ca, cs = (col("country"), col("university"), col("official website"),
                          col("admissions page", "admission page"), col("status"))
    country_default = path.stem.split("_")[0].capitalize()
    out = []
    for n, r in enumerate(rows[1:], start=2):
        name = r[cn] if cn is not None and cn < len(r) else None
        if not name:
            continue
        pages = str(r[ca] or "") if ca is not None and ca < len(r) else ""
        out.append(Uni(
            country=str(r[ci] or country_default) if ci is not None else country_default,
            name=str(name).strip(),
            website=str(r[cw] or "").strip() if cw is not None else "",
            admission_pages=[p for p in re.split(r"[\s,;]+", pages) if p.startswith("http")],
            status=str(r[cs] or "").strip() if cs is not None and cs < len(r) else "",
            row=n))
    return out


# --------------------------------------------------------------------------- page parsing
def page_text(p: Page) -> tuple[str, list[tuple[str, str]]]:
    """Visible text and (url, anchor text) links of a fetched page."""
    if p.is_pdf:
        try:
            reader = PdfReader(io.BytesIO(p.content))
            return "\n".join((pg.extract_text() or "") for pg in reader.pages[:25]), []
        except Exception:  # noqa: BLE001
            return "", []
    soup = BeautifulSoup(p.content, "lxml")
    links = []
    for a in soup.find_all("a", href=True):
        links.append((urljoin(p.final_url or p.url, a["href"].strip()), a.get_text(" ", strip=True)))
    for t in soup(["script", "style", "noscript", "svg", "iframe", "nav", "footer"]):
        t.decompose()
    for cell in soup.find_all(["td", "th"]):
        cell.insert_after(" | ")
    for t in soup.find_all(["p", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "dt", "dd",
                            "section", "article", "br", "table", "ul", "ol"]):
        t.insert_before("\n")
        t.insert_after("\n")
    lines = [re.sub(r"\s+", " ", ln).strip(" |") for ln in soup.get_text(" ").split("\n")]
    return "\n".join(ln for ln in lines if ln), links


def base_domain(host: str) -> str:
    host = host.lower().split(":")[0]
    parts = host.split(".")
    return ".".join(parts[-2:]) if len(parts) >= 2 else host


def rank_links(links: list[tuple[str, str]], domain: str, seen: set[str], today: date | None = None) -> list[str]:
    year = (today or date.today()).year
    scored = {}
    for url, text in links:
        url = url.split("#")[0]
        p = urlparse(url)
        if p.scheme not in ("http", "https") or base_domain(p.netloc) != domain or url in seen:
            continue
        blob = f"{url} {text}"
        if LINK_BAD.search(url) or LINK_BAD.search(text or "") or NOT_ADMISSION_PAGE.search(url):
            continue
        score = 3 * bool(LINK_STRONG.search(blob)) + bool(LINK_MEDIUM.search(blob))
        if url.lower().endswith(".pdf") and score:
            score += 1
        if re.search(f"{year}|{year + 1}", blob):
            score += 1
        if score >= 2:
            scored[url] = max(score, scored.get(url, 0))
    return sorted(scored, key=lambda u: -scored[u])


# --------------------------------------------------------------------------- checking
@dataclass
class Result:
    uni: Uni
    matches: list[tuple[str, Window]] = field(default_factory=list)
    problem: str = ""
    detail: str = ""
    pages_ok: int = 0
    blocked: list[str] = field(default_factory=list)
    robots_skipped: list[str] = field(default_factory=list)
    methods: set = field(default_factory=set)


class BrowserRouted:
    """Fetcher stand-in for the browser retry: pages on hosts the HTTP fetcher found blocked go
    through the browser (robots.txt still checked, one page at a time); other hosts use HTTP."""

    def __init__(self, fetcher: Fetcher, br, country: str):
        self.fetcher, self.br, self.country = fetcher, br, country
        self.blocked_hosts = fetcher.blocked_hosts

    def get(self, url: str, referer: str | None = None) -> Page:
        if urlparse(url).netloc.lower() not in self.fetcher.blocked_hosts:
            page = self.fetcher.get(url, referer)
            if not page.status.startswith("blocked"):
                return page
        allowed, why = self.fetcher.allowed(url)
        if not allowed:
            return Page(url, f"skipped: {why}")
        time.sleep(random.uniform(*self.fetcher.delay))
        return self.br.get(url, self.country)


def check(uni: Uni, fetcher: Fetcher, cfg: dict, today: date, extra: dict[str, Page] | None = None) -> Result:
    res = Result(uni)
    start = uni.admission_pages or ([uni.website] if uni.website else [])
    if not start:
        res.problem, res.detail = "No official website in sheet", "fill 'Official Website' or 'Admissions Page'"
        return res
    domain = base_domain(urlparse(start[0]).netloc)
    queue, seen, windows, errors = list(start), set(), [], []
    max_pages = int(cfg["MAX_PAGES"])
    while queue and len(seen) < max_pages:
        url = queue.pop(0)
        if url in seen:
            continue
        seen.add(url)
        page = (extra or {}).get(url) or fetcher.get(url, referer=start[0] if url != start[0] else None)
        if not page.ok and page.status.startswith("blocked"):
            page = wayback(url) or page
        if not page.ok:
            errors.append(f"{page.status}: {url}")
            if page.status.startswith("skipped:"):
                res.robots_skipped.append(url)
            if page.status.startswith("blocked"):
                res.blocked.append(url)
            continue
        res.pages_ok += 1
        res.methods.add(page.method.split(":")[0])
        text, links = page_text(page)
        page_url = page.final_url or url
        if url in start or not NOT_ADMISSION_PAGE.search(page_url):
            windows += find_windows(text, page_url, today)
        for link in rank_links(links, domain, seen | set(queue), today):
            if len(queue) < max_pages * 2:
                queue.append(link)
    dd, od = int(cfg["DEADLINE_DAYS"]), int(cfg["OPENED_DAYS"])
    uniq = {}
    for w in windows:
        st = w.status(today, dd, od)
        if st and st != "Opens soon" and (w.opens, w.closes) not in uniq:
            uniq[(w.opens, w.closes)] = (st, w)
    order = {"OPEN NOW": 0, "Deadline coming (opening date not stated)": 1,
             "Opened recently (closing date not stated)": 2}
    res.matches = sorted(uniq.values(), key=lambda sw: (order[sw[0]], sw[1].closes or date.max))[:6]
    if res.pages_ok == 0 and res.robots_skipped and not res.blocked:
        res.problem = "Not checked: site's robots.txt does not allow it"
        res.detail = "; ".join(errors[:3])
    elif res.pages_ok == 0:
        res.problem = "Blocked by bot protection" if res.blocked else "Could not load official pages"
        res.detail = "; ".join(errors[:3])
    elif not windows:
        res.problem = "No admission dates found on official pages"
        res.detail = f"checked {res.pages_ok} page(s); set a more specific 'Admissions Page' in the sheet"
    return res


def academic_year(today: date) -> str:
    y = today.year if today.month >= 6 else today.year - 1
    return f"{y}/{y + 1}"


def index_queries(country: str, root: str, today: date) -> list[str]:
    ay, y = academic_year(today), today.year
    if country.lower() == "germany":
        ws = f"Wintersemester {ay[:4]}/{ay[-2:]}" if today.month >= 4 and today.month < 10 else f"Sommersemester {y + (today.month >= 10)}"
        return [f"site:{root} Bewerbungsfrist {ws}", f"site:{root} Bewerbung Zulassung Fristen {y}",
                f"site:{root} application deadline {y}"]
    return [f"site:{root} immatricolazioni {ay} scadenza", f"site:{root} bando ammissione {ay}",
            f"site:{root} admission deadline {y}"]


def index_fallback(res: Result, fetcher: Fetcher, cfg: dict, today: date) -> Result:
    """For universities whose main site demands human verification: read the same university's
    pages through (1) the search engine's index of its official domain (title + text snippet)
    and (2) official subdomains that are not behind the challenge."""
    start = res.uni.admission_pages or [res.uni.website]
    root = base_domain(urlparse(start[0]).netloc)
    hits = []
    for q in index_queries(res.uni.country, root, today):
        for r in search(q):
            url = r.get("href", "")
            if base_domain(urlparse(url).netloc) == root and url not in {h["href"] for h in hits}:
                hits.append(r)
        time.sleep(0.8)
    windows, pages = [], 0
    for r in hits:
        if NOT_ADMISSION_PAGE.search(r["href"]):
            continue
        text = f"{r.get('title', '')}\n{r.get('body', '')}"
        if re.search(r"universit[àa] (?:degli studi )?di (\w+)", text, re.I) and not any(
                w.lower() in text.lower() for w in re.findall(r"[A-Za-z]{5,}", res.uni.name)
                if w.lower() not in ("university", "universita", "studies", "polytechnic")):
            continue
        found = find_windows(text, r["href"], today)
        for w in found:
            w.context = "[search index of official site] " + w.context
        windows += found
        pages += 1
    # official subdomains that answer normally (department / portal sites)
    for r in hits[:8]:
        host = urlparse(r["href"]).netloc.lower()
        if host in fetcher.blocked_hosts or NOT_ADMISSION_PAGE.search(r["href"]):
            continue
        page = fetcher.get(r["href"])
        if page.ok:
            text, _ = page_text(page)
            windows += find_windows(text, page.final_url or r["href"], today)
            pages += 1
            res.methods.add("official subdomain")
    if pages:
        res.methods.add("search index")
        res.pages_ok += pages
        dd, od = int(cfg["DEADLINE_DAYS"]), int(cfg["OPENED_DAYS"])
        uniq = {}
        for w in windows:
            st = w.status(today, dd, od)
            if st and st != "Opens soon" and (w.opens, w.closes) not in uniq:
                uniq[(w.opens, w.closes)] = (st, w)
        res.matches = sorted(uniq.values(), key=lambda sw: (sw[0] != "OPEN NOW", sw[1].closes or date.max))[:6]
        if res.matches or windows:
            res.problem, res.detail = "", ""
        else:
            res.problem = "Main site needs human verification; no dates in search index / subdomains"
            res.detail = f"read {pages} indexed page(s) / subdomain page(s) of {root}"
    else:
        res.problem = "Blocked by human verification (CAPTCHA); nothing indexed"
    return res


# --------------------------------------------------------------------------- report sheets
def _style(ws, widths):
    fill = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def write_matches(path: Path, results: list[Result], today: date):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Open admissions"
    ws.append(["Country", "University", "Match", "Opens", "Closes", "Days left", "Evidence (page text)",
               "Page", "Official Website", "Read via"])
    for r in results:
        for st, w in r.matches:
            left = (w.closes - today).days if w.closes else ""
            ws.append([r.uni.country, r.uni.name, st, w.opens.isoformat() if w.opens else "",
                       w.closes.isoformat() if w.closes else "", left, w.context[:400], w.url, r.uni.website,
                       ", ".join(sorted(m for m in r.methods if m)) or "direct"])
    _style(ws, [10, 34, 26, 12, 12, 9, 70, 50, 30, 22])
    wb.save(path)


def write_problems(path: Path, results: list[Result]):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scrape problems"
    ws.append(["Country", "University", "Problem", "Details", "Pages loaded", "Blocked URLs", "Official Website",
               "Admissions Page (sheet)"])
    for r in results:
        ws.append([r.uni.country, r.uni.name, r.problem, r.detail, r.pages_ok, "\n".join(r.blocked[:3]),
                   r.uni.website, "\n".join(r.uni.admission_pages)])
    _style(ws, [10, 34, 34, 60, 8, 45, 30, 45])
    wb.save(path)


# --------------------------------------------------------------------------- email
def send_email(cfg: dict, subject: str, body: str, attachments: list[Path]):
    user, pwd = cfg.get("GMAIL_ADDRESS"), cfg.get("GMAIL_APP_PASSWORD", "").replace(" ", "")
    to = [a.strip() for a in cfg.get("NOTIFY_TO", user or "").split(",") if a.strip()]
    if not user or not pwd or not to:
        raise RuntimeError("GMAIL_ADDRESS / GMAIL_APP_PASSWORD / NOTIFY_TO missing in config.env")
    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"] = user, ", ".join(to), subject
    msg.set_content(body)
    for p in attachments:
        msg.add_attachment(p.read_bytes(), maintype="application",
                           subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=p.name)
    for attempt in range(3):
        try:
            with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=ssl.create_default_context(), timeout=60) as s:
                s.login(user, pwd)
                s.send_message(msg)
            return
        except smtplib.SMTPAuthenticationError:
            raise
        except Exception:  # noqa: BLE001
            if attempt == 2:
                raise
            time.sleep(10 * (attempt + 1))


def email_body(matches: list[Result], problems: list[Result], checked: int, skipped: int, secs: float,
               today: date | None = None) -> str:
    lines = [f"Admission watch - {(today or date.today()):%A %d %B %Y}", "",
             f"Checked {checked} universities in {secs / 60:.1f} min ({skipped} skipped: status Applied/Skip).", ""]
    if matches:
        lines.append(f"ADMISSIONS OPEN / DEADLINE COMING ({len(matches)} universities):")
        for r in matches:
            st, w = r.matches[0]
            when = f"{w.opens or '?'} -> {w.closes or '?'}"
            lines.append(f"  - {r.uni.name} ({r.uni.country}): {st}, {when}   {w.url}")
        lines.append("")
    else:
        lines += ["No open admission windows found today.", ""]
    if problems:
        lines.append(f"{len(problems)} universities could not be checked, see scrape_problems sheet.")
    lines += ["", "Set Status = Applied in your sheet to stop getting a university in this report."]
    return "\n".join(lines)


# --------------------------------------------------------------------------- main
def verify(unis: list[Uni], cfg: dict, timeout: float):
    """Open every university's first page in a visible browser; a person passes any human check."""
    if not browser.available():
        sys.exit("patchright is not installed (see README)")
    if not os.environ.get("DISPLAY"):
        sys.exit("--verify needs a visible screen: run it from a desktop session (or ssh -X)")
    robots = RobotsPolicy() if cfg["RESPECT_ROBOTS"].lower() != "false" else None
    fetchers = {c: Fetcher(c, robots=robots) for c in {u.country for u in unis}}
    br = browser.StealthBrowser(headless=False)
    status = {}
    try:
        for u in unis:
            url = (u.admission_pages or [u.website])[0]
            if not url:
                continue
            allowed, why = fetchers[u.country].allowed(url)
            status[u.name] = br.manual_verify(url, u.country, timeout) if allowed else f"skipped: {why}"
            print(f"  {u.name[:45]:45} {status[u.name]}", flush=True)
    finally:
        br.close()
    ok = sum(s == "OK" for s in status.values())
    print(f"\n{ok}/{len(status)} reachable. Sessions saved in {browser.STATE_DIR}")


def main():
    ap = argparse.ArgumentParser(description="Nightly check of official admission pages")
    ap.add_argument("--dry-run", action="store_true", help="don't send email, print results")
    ap.add_argument("--keep", action="store_true", help="keep the report sheets in ./reports")
    ap.add_argument("--only", help="comma-separated parts of university names to check")
    ap.add_argument("--country", help="only this country")
    ap.add_argument("--today", help="pretend today is YYYY-MM-DD (testing)")
    ap.add_argument("--test-email", action="store_true", help="only send a test email and exit")
    ap.add_argument("--verify", action="store_true",
                    help="open each university's page in a visible browser so you can pass its human check")
    ap.add_argument("--verify-timeout", type=float, default=180, help="seconds to wait per site in --verify")
    args = ap.parse_args()
    if args.test_email:
        cfg = load_config()
        send_email(cfg, "[Admission Watch] test email", "Gmail sending works. The nightly report will look like this.", [])
        print("test email sent to", cfg.get("NOTIFY_TO") or cfg.get("GMAIL_ADDRESS"))
        return

    (HERE / "logs").mkdir(exist_ok=True)
    logging.basicConfig(level=logging.WARNING, format="%(asctime)s %(message)s", datefmt="%H:%M:%S",
                        handlers=[logging.StreamHandler(sys.stdout),
                                  logging.FileHandler(HERE / "logs" / f"run_{date.today()}.log", encoding="utf-8")])
    log.setLevel(logging.INFO)           # only our own messages; search libraries stay quiet
    for f in (HERE / "logs").glob("run_*.log"):          # keep two weeks of logs
        if time.time() - f.stat().st_mtime > 14 * 86400:
            f.unlink()
    cfg = load_config()
    today = date.fromisoformat(args.today) if args.today else date.today()
    t0 = time.time()

    unis: list[Uni] = []
    for name in cfg["SHEETS"].split(","):
        path = HERE / name.strip()
        if path.exists():
            unis += read_sheet(path)
        else:
            log.warning("sheet not found: %s", path)
    if args.country:
        unis = [u for u in unis if u.country.lower() == args.country.lower()]
    if args.only:
        keys = [k.strip().lower() for k in args.only.split(",")]
        unis = [u for u in unis if any(k in u.name.lower() for k in keys)]
    skipped = [u for u in unis if u.status.lower() in SKIP_STATUSES]
    todo = [u for u in unis if u.status.lower() not in SKIP_STATUSES]
    if args.verify:
        verify(todo, cfg, args.verify_timeout)
        return
    log.info("checking %d universities (%d skipped by Status)", len(todo), len(skipped))

    proxies = [p.strip() for p in cfg.get("PROXIES", "").split(",") if p.strip()]
    robots = RobotsPolicy() if cfg["RESPECT_ROBOTS"].lower() != "false" else None
    fetchers = {c: Fetcher(c, proxies, robots=robots) for c in {u.country for u in todo}}
    results: list[Result] = []
    with cf.ThreadPoolExecutor(int(cfg["WORKERS"])) as ex:
        futs = {ex.submit(check, u, fetchers[u.country], cfg, today): u for u in todo}
        for f in cf.as_completed(futs):
            u = futs[f]
            try:
                r = f.result()
            except Exception as e:  # noqa: BLE001
                r = Result(u, problem="Script error", detail=f"{type(e).__name__}: {e}")
            results.append(r)
            first = r.matches[0] if r.matches else None
            log.info("  %-45s %s", u.name[:45], f"{first[0]} {first[1].opens}->{first[1].closes}" if first
                     else (r.problem or "no open window"))

    # second chance for blocked / JS-only sites in a stealth headless browser (sequential)
    retry = [r for r in results if r.pages_ok == 0 and r.blocked]
    if retry and browser.available():
        log.info("retrying %d blocked universities in stealth browser", len(retry))
        br = browser.open_browser()
        log.info("browser engine: %s", type(br).__name__)
        try:
            for r in retry:
                url = r.blocked[0]
                allowed, why = fetchers[r.uni.country].allowed(url)
                if not allowed:
                    r.detail = f"browser retry skipped ({why}); " + r.detail
                    continue
                page = br.get(url, r.uni.country)
                if not page.ok:
                    hint = f' (run: admission_watch.py --verify --only "{r.uni.name}")' \
                        if "human-verification" in page.status else ""
                    r.detail = f"{page.status}{hint}; " + r.detail
                    log.info("  %-45s %s", r.uni.name[:45], page.status)
                    continue
                routed = BrowserRouted(fetchers[r.uni.country], br, r.uni.country)
                new = check(r.uni, routed, cfg, today, extra={url: page})
                new.methods.add("stealth-browser")
                results[results.index(r)] = new
                log.info("  %-45s %s", r.uni.name[:45], new.problem or f"{len(new.matches)} match(es)")
        finally:
            br.close()

    # third chance: the search index of the official domain + unprotected official subdomains
    still = [r for r in results if r.pages_ok == 0 and r.blocked]
    # the browser got past the wall but the pages it reached had no dates: the index may still have them
    thin = [r for r in results if "stealth-browser" in r.methods and not r.matches]
    if still or thin:
        log.info("reading %d still-blocked (+%d browser-reached, no dates) universities via search index / "
                 "official subdomains", len(still), len(thin))
        before = {id(r): (r.problem, r.detail) for r in thin}

        def fallback(r: Result) -> Result:
            r = index_fallback(r, fetchers[r.uni.country], cfg, today)
            if id(r) in before and not r.matches:       # keep the browser result's own problem text
                r.problem, r.detail = before[id(r)]
            return r
        with cf.ThreadPoolExecutor(4) as ex:
            done = list(ex.map(fallback, still + thin))
        for r in done:
            first = r.matches[0] if r.matches else None
            log.info("  %-45s %s", r.uni.name[:45], f"{first[0]} {first[1].opens}->{first[1].closes}"
                     if first else (r.problem or "checked, no open window"))

    matches = sorted([r for r in results if r.matches], key=lambda r: (r.uni.country, r.uni.name))
    problems = sorted([r for r in results if r.problem], key=lambda r: (r.uni.country, r.problem, r.uni.name))
    secs = time.time() - t0
    log.info("done in %.0fs: %d with open/closing admissions, %d problems", secs, len(matches), len(problems))

    outdir = Path(tempfile.mkdtemp(prefix="admission_watch_"))
    try:
        files = []
        stamp = today.isoformat()
        if matches:
            files.append(outdir / f"open_admissions_{stamp}.xlsx")
            write_matches(files[-1], matches, today)
        if problems:
            files.append(outdir / f"scrape_problems_{stamp}.xlsx")
            write_problems(files[-1], problems)
        body = email_body(matches, problems, len(todo), len(skipped), secs, today)
        subject = (f"[Admission Watch] {len(matches)} open admission(s), {len(problems)} problem(s) - {stamp}")
        if args.dry_run:
            print("\n" + subject + "\n" + body)
        elif matches or problems or cfg["SEND_WHEN_EMPTY"].lower() == "true":
            send_email(cfg, subject, body, files)
            log.info("email sent to %s", cfg.get("NOTIFY_TO") or cfg.get("GMAIL_ADDRESS"))
        if args.keep and files:
            keep = HERE / "reports"
            keep.mkdir(exist_ok=True)
            for p in files:
                shutil.copy(p, keep / p.name)
            log.info("kept copies in %s", keep)
    finally:
        shutil.rmtree(outdir, ignore_errors=True)       # no report files left behind


if __name__ == "__main__":
    main()
