#!/usr/bin/env python3
"""
University intake-date scraper.

Give it a country + university name. It will:
  1. find the official website (Hipolabs universities API -> web search fallback)
  2. search the web (free DuckDuckGo/Bing/Brave meta-search via `ddgs`) in English
     AND in the country's local language
  3. fetch pages with a real-browser TLS fingerprint (curl_cffi); if a page is
     bot-protected it retries with cloudscraper, then the Wayback Machine copy,
     and otherwise simply moves on to third-party sites that republish the dates
  4. read HTML and PDF pages, auto-detect language and translate to English
  5. extract intake START dates, application OPEN dates and application DEADLINES
  6. group them into intakes and label: Previous, Current, Next
     (single-intake uni -> 3 rows, multi-intake uni -> Previous-2, Previous-1, Current, Next)
  7. write a formatted Excel workbook (Summary / Intakes / Evidence / Sources)

If no previous intake was found on the first pass, a second pass searches
specifically for last year's dates and archived (Wayback) copies of official pages.

Usage:
  python intake_scraper.py -c Germany -u "Technical University of Munich"
  python intake_scraper.py -i universities.csv -o intake_dates.xlsx
  python intake_scraper.py            # interactive prompt
"""
from __future__ import annotations

import argparse
import os
import concurrent.futures as cf
import io
import logging
import re
import sys
import threading
import time
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path
from urllib.parse import urlparse

import cloudscraper
import pandas as pd
import requests
from bs4 import BeautifulSoup
from curl_cffi import requests as creq
from ddgs import DDGS
from deep_translator import GoogleTranslator
from langdetect import DetectorFactory, detect
from pypdf import PdfReader

from registry import DEFAULT_REGISTRY, fold, load_registry, match_country, parse_selection

DetectorFactory.seed = 0
log = logging.getLogger("intake")
logging.getLogger("pypdf").setLevel(logging.ERROR)

TODAY = date.today()

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #
COUNTRY_LANG = {
    "germany": "de", "austria": "de", "switzerland": ("de", "fr", "it"), "liechtenstein": "de",
    "france": "fr", "belgium": ("nl", "fr"), "luxembourg": ("fr", "de"), "morocco": "fr", "tunisia": "fr",
    "malta": "en",
    "senegal": "fr", "spain": "es", "mexico": "es", "argentina": "es", "chile": "es",
    "colombia": "es", "peru": "es", "venezuela": "es", "ecuador": "es", "uruguay": "es",
    "italy": "it", "netherlands": "nl", "portugal": "pt", "brazil": "pt", "sweden": "sv",
    "norway": "no", "denmark": "da", "finland": "fi", "iceland": "is", "poland": "pl",
    "czech republic": "cs", "czechia": "cs", "slovakia": "sk", "hungary": "hu",
    "romania": "ro", "bulgaria": "bg", "greece": "el", "cyprus": "el", "croatia": "hr",
    "serbia": "sr", "slovenia": "sl", "estonia": "et", "latvia": "lv", "lithuania": "lt",
    "turkey": "tr", "turkiye": "tr", "russia": "ru", "ukraine": "uk", "belarus": "ru",
    "kazakhstan": "ru", "japan": "ja", "china": "zh-CN", "taiwan": "zh-TW",
    "hong kong": "zh-TW", "south korea": "ko", "korea": "ko", "indonesia": "id",
    "vietnam": "vi", "thailand": "th", "malaysia": "ms", "saudi arabia": "ar",
    "united arab emirates": "ar", "uae": "ar", "egypt": "ar", "qatar": "ar",
    "jordan": "ar", "kuwait": "ar", "oman": "ar", "bahrain": "ar", "lebanon": "ar",
    "iran": "fa", "israel": "iw", "georgia": "ka", "armenia": "hy",
}
NAME_LANG_HINTS = [  # the institution's own name tells which local language it uses
    (re.compile(r"universit[ée]|[ée]cole|facult[ée]|\bde la\b|\bdes\b", re.I), "fr"),
    (re.compile(r"universit[äa]t|hochschule|akademie|technische|\bf[üu]r\b", re.I), "de"),
    (re.compile(r"universit[àa]\b|istituto|scuola|politecnico|\bdegli\b|\bdella\b", re.I), "it"),
    (re.compile(r"universiteit|hogeschool", re.I), "nl"),
    (re.compile(r"universidad|universitat\b|polit[ée]cnica", re.I), "es"),
]
HIPOLABS_COUNTRY = {"czechia": "Czech Republic", "turkiye": "Turkey", "korea": "Korea, Republic of",
                    "south korea": "Korea, Republic of", "usa": "United States", "us": "United States",
                    "uk": "United Kingdom"}


def local_languages(country: str, names: list[str]) -> list[str]:
    """Up to two non-English search languages for this institution."""
    base = COUNTRY_LANG.get(country.lower(), "en")
    langs = list(base) if isinstance(base, tuple) else [base]
    for rx, lang in NAME_LANG_HINTS:
        if lang in langs and any(rx.search(n) for n in names):
            langs.remove(lang)
            langs.insert(0, lang)
            break
    return [lang for lang in langs if lang != "en"][:2]


MDY_COUNTRIES = {"united states", "usa", "us", "united states of america", "philippines"}

SKIP_DOMAINS = (
    "facebook.com", "instagram.com", "youtube.com", "linkedin.com", "twitter.com", "x.com",
    "tiktok.com", "pinterest.", "quora.com", "reddit.com", "wikipedia.org", "wikidata.org",
    "glassdoor.", "indeed.", "amazon.", "apple.com", "google.com", "maps.", "tripadvisor.",
)
AGGREGATORS = (
    "topuniversities.com", "shiksha.com", "leverageedu.com", "yocket.com", "universityliving.com",
    "idp.com", "collegedunia.com", "edvoy.com", "leapscholar.com", "mastersportal.com",
    "bachelorsportal.com", "studyportals.com", "hotcoursesabroad.com", "timeshighereducation.com",
    "daad.de", "uni-assist.de", "mygermanuniversity.com", "study-in-germany.de", "upgrad.com",
    "gotouniversity.com", "unischolars.com", "amberstudent.com", "studyinternational.com",
    "educations.com", "univariety.com", "getmyuni.com", "careers360.com", "studying-in-",
    "kcl-edu", "abroadvice.com", "geeksforgeeks.org", "unirank.org", "4icu.org",
)
NAME_STOPWORDS = {
    "university", "universitat", "universität", "universite", "université", "universidad",
    "universita", "università", "of", "the", "and", "de", "di", "du", "des", "la", "le", "für",
    "fur", "institute", "college", "school", "national", "state", "technical", "technology",
    "sciences", "science", "applied", "for", "at", "in", "a", "an", "&",
}

FOREIGN_CAMPUS_WORDS = set(COUNTRY_LANG) | {
    "australia", "india", "united kingdom", "uk", "canada", "singapore", "new zealand", "ireland",
    "south africa", "pakistan", "bangladesh", "sri lanka", "nigeria", "kenya", "ghana", "dubai",
}

ADMIN_PATH = re.compile(r"census|exam|timetable|graduation|convocation|fees?\b|payment|results|withdraw|holiday")

BROWSER_HEADERS = {"Accept-Language": "en-US,en;q=0.9"}
BLOCK_MARKERS = (
    "just a moment", "cf-browser-verification", "cf_chl_", "attention required",
    "are you a robot", "verify you are human", "access denied", "incapsula", "datadome",
    "px-captcha", "request unsuccessful", "enable javascript and cookies", "captcha",
    "bot detection", "unusual traffic",
)

# --------------------------------------------------------------------------- #
# Date patterns (applied to English / translated text)
# --------------------------------------------------------------------------- #
_MONTH_NAMES = [
    ("january", "jan"), ("february", "feb"), ("march", "mar"), ("april", "apr"), ("may",),
    ("june", "jun"), ("july", "jul"), ("august", "aug"), ("september", "sept", "sep"),
    ("october", "oct"), ("november", "nov"), ("december", "dec"),
]
MONTHS = {n: i for i, names in enumerate(_MONTH_NAMES, 1) for n in names}
MON = (r"(jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?"
       r"|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)\.?(?![a-z])")
YEAR = r"(20\d{2})(?!\d)"
ORD = r"(?:st|nd|rd|th)?"
RSEP = r"(?:[-–—]|to|until|till|and|through)"

DATE_PATTERNS = [  # (name, regex) — order matters, earlier patterns win overlapping spans
    ("range_dmy", re.compile(rf"\b(\d{{1,2}}){ORD}\s*{RSEP}\s*(\d{{1,2}}){ORD}\s+(?:of\s+)?{MON}\s*,?\s*{YEAR}", re.I)),
    ("range_dm", re.compile(rf"\b(\d{{1,2}}){ORD}\s*{RSEP}\s*(\d{{1,2}}){ORD}\s+(?:of\s+)?{MON}", re.I)),
    ("dmy", re.compile(rf"\b(\d{{1,2}}){ORD}\s+(?:of\s+)?{MON}\s*,?\s*{YEAR}", re.I)),
    ("mdy", re.compile(rf"\b{MON}\s+(\d{{1,2}}){ORD}\s*,?\s*{YEAR}", re.I)),
    ("iso", re.compile(r"\b(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})\b")),
    ("num", re.compile(r"\b(\d{1,2})[./-](\d{1,2})[./-](20\d{2})\b")),
    ("my", re.compile(rf"\b{MON}\s*,?\s*{YEAR}", re.I)),
    ("dm", re.compile(rf"\b(\d{{1,2}}){ORD}\s+(?:of\s+)?{MON}", re.I)),
    ("md", re.compile(rf"\b{MON}\s+(\d{{1,2}}){ORD}\b(?!\s*[,.]?\s*\d)", re.I)),
]
RANGE_GAP = re.compile(r"^\s*(?:-|–|—|to|until|till|through|thru|and|bis)\s*$", re.I)
YEAR_RE = re.compile(r"(?<!\d)(20\d{2})(?!\d)")

# Keyword classes used to decide what a date means.
KW = {
    "start": re.compile(
        r"\b(intakes?|(?:semester|term|trimester|quarter|session|classes|lectures?|teaching|courses?|studies"
        r"|programm?e|academic year|instruction|lecture period)\s+(?:begins?|starts?|commences?)"
        r"|start(?:ing)? dates?|course start|programme start|program start|semester start|lectures? start"
        r"|first day of (?:classes|term|semester|lectures|instruction|school)"
        r"|(?:start|beginning|begin|commencement) of (?:the )?(?:semester|term|trimester|lectures"
        r"|classes|studies|instruction|academic year|lecture period|programme|program|course)"
        r"|lecture period|teaching period|enrol?ment (?:day|week)|welcome week|entry)\b", re.I),
    "deadline": re.compile(
        r"\b(deadlines?|apply by|applications? (?:close|closes|closing|due|must be (?:received|submitted))"
        r"|closing dates?|last (?:date|day) (?:to|for) (?:apply|application|applications|submission|submit)"
        r"|last date|submission (?:deadline|date)|application period ends|apply (?:until|before)"
        r"|applications? (?:\w+ ){0,4}(?:until|by|before|no later than)|no later than|due date"
        r"|(?:and|applications?|portal|window)\s+close[sd]?|closes? on|closing)\b", re.I),
    "open": re.compile(
        r"\b(applications? (?:open|opens|start|starts|begin|begins)|application (?:period|window|cycle|portal)"
        r" (?:opens|begins|starts)|opens? for applications|apply from|application start"
        r"|start of (?:the )?application|applications? (?:can|may) be submitted (?:from|as of)"
        r"|opening dates?)\b", re.I),
}
KW_APPL = re.compile(r"\b(appl(?:y|ication|icants?)|admission|submission|register|registration)\b", re.I)
KW_EXCL = re.compile(
    r"\b(published|updated|posted|last modified|copyright|webinar|open days?|open house|exams?"
    r"|examinations?|results?|graduation|convocation|holidays?|break|recess|fee payment|payment"
    r"|tuition|withdraw\w*|refund|add/drop|census|re-?registration|re-?enrol\w*|scholarships?"
    r"|housing|accommodation|visa|news|events?|conference|workshop|ends?|end of|closes? for"
    r"|reading week|vacation|thesis|defen[cs]e|grades?|exchange|erasmus|learning agreements?"
    r"|nominations?|incoming|outgoing|matriculation|enrol?ment deadline|admitted students|waitlist"
    r"|enrol+ in|register for|course (?:selection|registration)|courses? in|drop|add courses?"
    r"|course enrol\w*|enrol\w* (?:period|day|opens|appointments?)|non-enrol\w*|registration for"
    r"|course changes|resumes?|continues?|continuing|returning students)\b|©", re.I)
SNIPPET_PREFIX = re.compile(
    r"^\s*(?:\d+\s+(?:minutes?|hours?|days?|weeks?|months?|years?)\s+ago|[A-Z][a-z]{2,8}\.?\s+\d{1,2},\s+\d{4}"
    r"|\d{1,2}\s+[A-Z][a-z]{2,8}\.?\s+\d{4}|\d{1,2}-[A-Z][a-z]{2}-\d{4}|\d{4}-\d{2}-\d{2})\s*[-·—–]\s*")
SNIPPET_MID = re.compile(  # "... Sep 14, 2026 · We release ..." (crawl/publish date inside a snippet)
    r"(?:[A-Z][a-z]{2,8}\.?\s+\d{1,2},\s+\d{4}|\d{1,2}\s+[A-Z][a-z]{2,8}\.?\s+\d{4}|\d{1,2}-[A-Z][a-z]{2}-\d{4})\s*·")
KW_INTAKE_WORDS = re.compile(r"\b(intake|admission|entry|appl\w+|enrol\w*|cohort|new students|freshers?)\b", re.I)

# Intake-pattern statements: "two intakes: February and July", "Fall and Spring intakes"
_TERM = (r"(?:jan(?:uary)?|feb(?:ruary)?|march|april|may|june|july|aug(?:ust)?|sept?(?:ember)?|oct(?:ober)?"
         r"|nov(?:ember)?|dec(?:ember)?|fall|autumn|spring|summer|winter)")
_TERM_LIST = rf"({_TERM}(?:\s*(?:,|and|&|or|/)\s*{_TERM})+)"
PATTERN_RES = [
    re.compile(rf"\b(?:intakes?|intake (?:months?|periods?|sessions?)|entry points?|admission (?:cycles?|sessions?|intakes?))"
               rf"\b[^.\n]{{0,40}}?\b{_TERM_LIST}\b", re.I),
    re.compile(rf"\b{_TERM_LIST}\s+(?:intakes?|entry|admission (?:cycles?|sessions?))\b", re.I),
]
SEASON_MONTH = {"fall": 9, "autumn": 9, "spring": 2, "summer": 5, "winter": 1}

# Hints: which intake does a deadline refer to?  -> (year | None, month)
HINTS = [
    (re.compile(r"\bwinter\s*(?:semester|sem\.?)\s*(20\d{2})?", re.I), 10),
    (re.compile(r"\bsummer\s*(?:semester|sem\.?)\s*(20\d{2})?", re.I), 4),
    (re.compile(r"\b(?:fall|autumn)\s*(?:semester|term|intake|session|quarter|entry|admissions?|cohort)?\s*,?\s*(20\d{2})\b", re.I), 9),
    (re.compile(r"\b(?:fall|autumn)\s+(?:semester|term|intake|session|quarter|entry|admissions?|cohort)()", re.I), 9),
    (re.compile(r"\bspring\s*(?:semester|term|intake|session|quarter|entry|admissions?|cohort)?\s*,?\s*(20\d{2})\b", re.I), 2),
    (re.compile(r"\bspring\s+(?:semester|term|intake|session|quarter|entry|admissions?|cohort)()", re.I), 2),
    (re.compile(r"\bwinter\s+(?:term|quarter|session|intake|entry)\s*(20\d{2})?", re.I), "winter-term"),
    (re.compile(r"\bsummer\s+(?:term|session|intake|quarter|entry)\s*(20\d{2})?", re.I), "summer-term"),
]
HINT_MONTH_INTAKE = re.compile(
    rf"(?:{MON}\s*,?\s*(?:(20\d{{2}})\s+)?(?:intake|entry|admissions?|start|cohort))"
    rf"|(?:(?:intake|entry|cohort)\s*(?:for|in|of)?\s*:?\s*{MON}\s*,?\s*(20\d{{2}})?)", re.I)


# --------------------------------------------------------------------------- #
# Data classes
# --------------------------------------------------------------------------- #
@dataclass
class Candidate:
    kind: str               # start | deadline | open
    date: date
    precision: str          # day | month
    score: float
    url: str
    source_type: str        # official | official-archive | third-party | search-snippet
    context: str
    hint: tuple | None = None
    translated: bool = False
    virtual: bool = False   # start month inferred from an intake label (e.g. "Fall 2026")


@dataclass
class Intake:
    start: date
    start_precision: str
    start_score: float
    cands: list = field(default_factory=list)
    deadlines: list = field(default_factory=list)
    opens: list = field(default_factory=list)
    inferred: bool = False
    found: bool = False
    slot: date | None = None   # nominal month of this intake in the learned pattern

    @property
    def score(self):
        return self.start_score + sum(c.score for c in self.deadlines) * 0.5


# --------------------------------------------------------------------------- #
# Networking
# --------------------------------------------------------------------------- #
_scraper = cloudscraper.create_scraper()


def looks_blocked(status: int, html: str) -> bool:
    if status in (401, 403, 406, 429, 503):
        return True
    low = html[:80000].lower()
    return len(html) < 60000 and any(m in low for m in BLOCK_MARKERS)


def http_get(url: str, timeout: int = 20):
    """Return (status_str, content_bytes, content_type, final_url)."""
    attempts = (
        ("curl_cffi", lambda: creq.get(url, impersonate="chrome", timeout=timeout,
                                       headers=BROWSER_HEADERS, allow_redirects=True)),
        ("cloudscraper", lambda: _scraper.get(url, timeout=timeout, headers=BROWSER_HEADERS)),
    )
    last = "error"
    for name, fn in attempts:
        try:
            r = fn()
            ctype = r.headers.get("content-type", "")
            body = r.content
            text = body[:200000].decode("utf-8", "ignore") if "pdf" not in ctype else ""
            if looks_blocked(r.status_code, text):
                last = f"blocked({r.status_code})"
                continue
            if r.status_code >= 400:
                return f"http {r.status_code}", b"", ctype, url
            return "ok", body, ctype, str(r.url)
        except Exception as e:  # noqa: BLE001
            last = f"error: {type(e).__name__}"
    return last, b"", "", url


def wayback_get(url: str, timestamp: date | None = None):
    """Fetch the closest Wayback Machine snapshot (raw, without the archive toolbar)."""
    ts = (timestamp or TODAY).strftime("%Y%m%d")
    try:
        r = requests.get("https://archive.org/wayback/available",
                         params={"url": url, "timestamp": ts}, timeout=20)
        snap = r.json().get("archived_snapshots", {}).get("closest")
        if not snap or not snap.get("available"):
            return None
        raw = re.sub(r"/web/(\d+)/", r"/web/\1id_/", snap["url"], count=1)
        status, body, ctype, _ = http_get(raw, timeout=30)
        if status == "ok":
            return body, ctype, raw
    except Exception:  # noqa: BLE001
        pass
    return None


SEARCH_BACKENDS = ("auto", "yahoo", "duckduckgo", "brave", "mojeek")


def search(query: str, max_results: int = 10) -> list[dict]:
    """Free meta-search; walks through engines until one answers (others may rate-limit)."""
    for attempt, backend in enumerate(SEARCH_BACKENDS):
        try:
            with DDGS() as d:
                res = list(d.text(query, region="wt-wt", safesearch="off",
                                  max_results=max_results, backend=backend))
            if res:
                return res
        except Exception as e:  # noqa: BLE001
            log.debug("search %s failed (%s) %s", backend, e, query)
        time.sleep(1 + attempt)
    return []


def domain_of(url: str) -> str:
    host = urlparse(url).netloc.lower().split(":")[0]
    return host[4:] if host.startswith("www.") else host


def find_official_domain(names: list[str] | str, country: str) -> str | None:
    names = [names] if isinstance(names, str) else names
    keys = {frozenset(name_tokens(n)) for n in names}
    hip_country = HIPOLABS_COUNTRY.get(country.lower(), country)
    for name in names:
        for params in ({"name": name, "country": hip_country}, {"name": name}):
            try:
                r = requests.get("http://universities.hipolabs.com/search", params=params, timeout=15)
                data = r.json()
            except Exception:  # noqa: BLE001
                continue
            # a fuzzy hit like "Medical University of Graz" for "University of Graz" is not good enough
            for d in data:
                if frozenset(name_tokens(d.get("name", ""))) in keys and d.get("domains"):
                    return d["domains"][0].lower().removeprefix("www.")
    tokens = [t for n in names for t in name_tokens(n)]
    for res in search(f"{names[0]} {country} official website", 8):
        dom = domain_of(res.get("href", ""))
        if not dom or any(s in dom for s in SKIP_DOMAINS + AGGREGATORS):
            continue
        blob = fold(dom + " " + res.get("title", ""))
        if any(t in blob for t in tokens) or re.search(r"\.(edu|ac\.|uni-)", dom):
            return dom
    return None


# --------------------------------------------------------------------------- #
# Text extraction + translation
# --------------------------------------------------------------------------- #
BLOCK_TAGS = ["p", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "dt", "dd",
              "section", "article", "table", "br", "ul", "ol", "header", "footer", "td", "th"]


def html_to_text(html: bytes) -> tuple[str, str]:
    soup = BeautifulSoup(html, "lxml")
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    for t in soup(["script", "style", "noscript", "svg", "iframe", "form"]):
        t.decompose()
    for cell in soup.find_all(["td", "th"]):
        cell.insert_after(" | ")
    for t in soup.find_all(BLOCK_TAGS):
        if t.name in ("td", "th"):
            continue
        t.insert_before("\n")
        t.insert_after("\n")
    raw = soup.get_text(" ")
    lines = [re.sub(r"\s+", " ", ln).strip(" |") for ln in raw.split("\n")]
    return title, "\n".join(ln for ln in lines if ln)


def pdf_to_text(data: bytes, max_pages: int = 20) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages[:max_pages])
    except Exception:  # noqa: BLE001
        return ""


def detect_lang(text: str) -> str:
    try:
        return detect(text[:4000])
    except Exception:  # noqa: BLE001
        return "en"


_tr_cache: dict[str, str] = {}


def translate(text: str, target: str = "en", source: str = "auto") -> str:
    key = f"{source}>{target}:{text}"
    if key in _tr_cache:
        return _tr_cache[key]
    out = text
    for attempt in range(2):
        try:
            out = GoogleTranslator(source=source, target=target).translate(text) or text
            break
        except Exception:  # noqa: BLE001
            time.sleep(1.5)
    _tr_cache[key] = out
    return out


def translate_relevant(text: str, max_chars: int = 14000) -> str:
    """Translate only lines likely to hold dates (plus the line before each, for labels)."""
    lines = text.split("\n")
    has_date = re.compile(r"(20\d{2}|\b\d{1,2}[./]\d{1,2}[./]?)")
    keep = set()
    for i, ln in enumerate(lines):
        if has_date.search(ln) and len(ln) < 600:
            keep.update({i - 1, i})
    picked = [lines[i] for i in sorted(keep) if 0 <= i < len(lines)]
    chunks, cur = [], ""
    total = 0
    for ln in picked:
        if total > max_chars:
            break
        if len(cur) + len(ln) + 1 > 4500:
            chunks.append(cur)
            cur = ""
        cur += ln + "\n"
        total += len(ln) + 1
    if cur:
        chunks.append(cur)
    return "\n".join(translate(c) for c in chunks)


YEAR_CELL = re.compile(r"^(20\d{2})(?:\s*[/–-]\s*(?:20)?\d{2})?$")


def annotate_table_years(text: str) -> str:
    """Calendar tables often put the year only in the header ("2026 | 2027 | 2028").
    Append each column's year to cells below it that have a day+month but no year."""
    out, header, age = [], None, 0
    for line in text.split("\n"):
        cells = [c.strip() for c in line.split(" | ")]
        if len(cells) >= 2 and all(YEAR_CELL.match(c) for c in cells if c) and sum(1 for c in cells if c) >= 2:
            header, age = [YEAR_CELL.match(c).group(1) for c in cells if c], 0
            out.append(line)
            continue
        age += 1
        if header and age <= 40 and len(cells) >= len(header):
            cols = cells[-len(header):]  # align from the right (first cell is often a row label)
            for i, c in enumerate(cols):
                if c and not YEAR_RE.search(c) and re.search(r"\d", c) and re.search(MON, c, re.I):
                    cols[i] = f"{c} {header[i]}"
            line = " | ".join(cells[:-len(header)] + cols)
        out.append(line)
    return "\n".join(out)


def name_tokens(uni: str) -> list[str]:
    stop = {fold(w) for w in NAME_STOPWORDS}
    toks = [t for t in re.findall(r"[\w'-]+", fold(uni)) if t not in stop and len(t) > 2]
    return toks or [fold(uni)]


def mentions_university(text: str, names: list[str] | str, acronym: str) -> bool:
    """True if the text names the institution (any alias, accents ignored) or its acronym."""
    low = fold(text)
    for name in ([names] if isinstance(names, str) else names):
        if all(t in low for t in name_tokens(name)):
            return True
    return bool(acronym) and len(acronym) >= 3 and re.search(rf"\b{re.escape(acronym)}\b", text) is not None


# --------------------------------------------------------------------------- #
# Date extraction
# --------------------------------------------------------------------------- #
def _mk_date(y, m, d):
    try:
        return date(int(y), int(m), int(d))
    except (ValueError, TypeError):
        return None


ACAD_YEAR_RE = re.compile(r"(?<!\d)(20\d{2})\s*[-/–]\s*(?:20)?(\d{2})(?!\d)")


def infer_year(text: str, s: int, e: int, month: int | None, fallback_year):
    """Year for a date written without one: nearby year > academic year (2026/27) > page's dominant year."""
    after = YEAR_RE.search(text[e:e + 45])
    if after:
        return after.group(1)
    near = text[max(0, s - 400):s]
    acad = [m for m in ACAD_YEAR_RE.finditer(near) if int(m.group(2)) == (int(m.group(1)) + 1) % 100]
    if acad and month:
        y1 = int(acad[-1].group(1))
        return y1 if month >= 8 else y1 + 1
    before = list(YEAR_RE.finditer(text[max(0, s - 80):s]))
    if before:
        return before[-1].group(1)
    return fallback_year


def find_dates(text: str, mdy: bool, fallback_year: int | None):
    """Yield (start, end, date, precision, range_partner_index|None)."""
    taken: list[tuple[int, int]] = []
    found = []

    def free(s, e):
        return all(e <= a or s >= b for a, b in taken)

    for name, rx in DATE_PATTERNS:
        for m in rx.finditer(text):
            s, e = m.span()
            if not free(s, e):
                continue
            g = m.groups()
            items = []
            if name in ("range_dmy", "range_dm"):
                mo = MONTHS.get(g[2].lower().rstrip("."))
                y = g[3] if name == "range_dmy" else infer_year(text, s, e, mo, fallback_year)
                if y is None:
                    continue
                items = [(_mk_date(y, mo, g[0]), "day"), (_mk_date(y, mo, g[1]), "day")]
            elif name == "dmy":
                items = [(_mk_date(g[2], MONTHS.get(g[1].lower().rstrip(".")), g[0]), "day")]
            elif name == "mdy":
                items = [(_mk_date(g[2], MONTHS.get(g[0].lower().rstrip(".")), g[1]), "day")]
            elif name == "iso":
                items = [(_mk_date(g[0], g[1], g[2]), "day")]
            elif name == "num":
                a, b, y = int(g[0]), int(g[1]), g[2]
                if a > 12:
                    dd, mm = a, b
                elif b > 12:
                    dd, mm = b, a
                else:
                    dd, mm = (b, a) if mdy else (a, b)
                items = [(_mk_date(y, mm, dd), "day")]
            elif name == "my":
                items = [(_mk_date(g[1], MONTHS.get(g[0].lower().rstrip(".")), 1), "month")]
            elif name in ("dm", "md"):
                day, mon = (g[0], g[1]) if name == "dm" else (g[1], g[0])
                mo = MONTHS.get(mon.lower().rstrip("."))
                y = infer_year(text, s, e, mo, fallback_year)
                if y is None:
                    continue
                items = [(_mk_date(y, mo, day), "day-noyear")]
            items = [(d, p) for d, p in items if d]
            if not items:
                continue
            taken.append((s, e))
            if len(items) == 2:
                found.append([s, e, items[0][0], "day", "range-a"])
                found.append([s, e, items[1][0], "day", "range-b"])
            else:
                found.append([s, e, items[0][0], items[0][1], None])
    found.sort(key=lambda x: (x[0], x[4] or ""))
    # detect "DATE to DATE" ranges
    for i in range(len(found) - 1):
        a, b = found[i], found[i + 1]
        if a[4] is None and b[4] is None and RANGE_GAP.match(text[a[1]:b[0]] or "x"):
            a[4], b[4] = "range-a", "range-b"
    return found


def _kw_score(before: str, after: str, rx: re.Pattern):
    """Proximity score of the nearest keyword match (0..1) and its absolute pos in `before`."""
    best, pos = 0.0, None
    for m in rx.finditer(before):
        dist = len(before) - m.end()
        sc = 1 / (1 + dist / 40)
        if "\n" in before[m.end():]:
            sc *= 0.6
        if sc > best:
            best, pos = sc, m.start()
    for m in rx.finditer(after):
        dist = m.start()
        sc = 0.7 / (1 + dist / 40)
        if "\n" in after[:m.start()]:
            sc *= 0.5
        if sc > best:
            best, pos = sc, None
    return best, pos


def extract_hint(ctx_before: str, ctx_after: str, na: bool = False):
    ctx_after = ctx_after.split("\n")[0]  # the next line usually belongs to another intake
    best, best_dist = None, 10 ** 6
    for side, txt in (("b", ctx_before), ("a", ctx_after)):
        for rx, month in HINTS:
            for m in rx.finditer(txt):
                y = m.group(1) if m.lastindex else None
                if month == "winter-term":  # North America: January; Europe: winter semester (Oct)
                    month = 1 if na else 10
                elif month == "summer-term":
                    month = 5 if na else 4
                dist = (len(txt) - m.end()) if side == "b" else m.start() + 30
                if dist < best_dist:
                    best, best_dist = ((int(y) if y else None), month), dist
        for m in HINT_MONTH_INTAKE.finditer(txt):
            g = m.groups()
            mon, y = (g[0], g[1]) if g[0] else (g[2], g[3])
            dist = (len(txt) - m.end()) if side == "b" else m.start() + 30
            if mon and dist < best_dist:
                best, best_dist = ((int(y) if y else None), MONTHS[mon.lower().rstrip(".")]), dist
    return best


def extract_pattern(text: str, url: str, source_type: str, weight: float, na: bool) -> list[Candidate]:
    """Month-of-year votes from sentences that describe the intake pattern (no year)."""
    out = []
    for rx in PATTERN_RES:
        for m in rx.finditer(text):
            named, seasons = set(), set()
            for tok in re.findall(_TERM, m.group(1), re.I):
                tok = tok.lower()
                if tok in SEASON_MONTH:
                    mo = SEASON_MONTH[tok]
                    if not na and tok == "winter":
                        mo = 10  # European "winter semester"
                    elif not na and tok == "summer":
                        mo = 4
                    seasons.add(mo)
                elif tok[:3] in MONTHS:
                    named.add(MONTHS[tok[:3]])
            # "January or Spring intake" names one intake twice
            months = named | {s for s in seasons if not any(_month_close(s, n, 1) for n in named)}
            ctx = text[max(0, m.start() - 60):m.end() + 40].replace("\n", " ¶ ")
            for mo in months:
                out.append(Candidate("pattern", date(2000, mo, 1), "month", round(weight, 3),
                                     url, source_type, ctx))
    return out


def extract_candidates(text: str, url: str, source_type: str, weight: float,
                       mdy: bool, translated: bool, na: bool = False) -> list[Candidate]:
    out = extract_pattern(text, url, source_type, weight, na)
    years = [int(y) for y in YEAR_RE.findall(text) if TODAY.year - 3 <= int(y) <= TODAY.year + 2]
    dominant = Counter(years).most_common(1)[0][0] if years else None
    for s, e, d, prec, rng in find_dates(text, mdy, dominant):
        if not (TODAY.year - 3 <= d.year <= TODAY.year + 2) or abs((d - TODAY).days) <= 1:
            continue  # out of range, or "today" (page-generated / crawl date)
        before = text[max(0, s - 140):s]
        after = text[e:e + 70]
        scores = {k: _kw_score(before, after, rx) for k, rx in KW.items()}
        kind, (sc, kpos) = max(scores.items(), key=lambda kv: kv[1][0])
        if sc < 0.22:
            continue
        if prec == "month" and re.match(r"\s*(?:intakes?|entry|cohort|semester|term|session)\b", after, re.I):
            kind, sc = "start", max(sc, 0.8)  # "September 2026 intake" is a label, never a deadline
        # "Application deadline for Winter Intake: 31 May" -> the intake word only labels a deadline
        if kind == "start" and kpos is not None and re.match(r"(?:intakes?|entry)\b", before[kpos:], re.I):
            other = max(("deadline", "open"), key=lambda k: scores[k][0])
            if scores[other][0] >= 0.3 and (scores[other][1] is None or scores[other][1] >= kpos - 60):
                kind, (sc, kpos) = other, scores[other]
        excl, _ = _kw_score(before, after, KW_EXCL)
        if excl > sc:
            continue
        if kpos is not None and KW_EXCL.search(before[max(0, kpos - 35):kpos + 25]):
            if not (kind == "deadline" and re.search(r"application", before[max(0, kpos - 35):], re.I)):
                continue
        # "1 June - 15 July 2026" with application context => open .. deadline
        if rng and KW_APPL.search(before[-(45 if rng == "range-a" else 80):].split("\n")[-1]):
            kind = "open" if rng == "range-a" else "deadline"
        elif rng == "range-b" and kind == "start":
            continue  # end of a semester range
        elif rng == "range-a" and kind == "deadline":
            continue  # "until" range: keep only the closing date
        p = "day" if prec.startswith("day") else "month"
        if kind == "start" and KW_INTAKE_WORDS.search(before[-80:] + after[:40]):
            sc *= 1.5
        conf = weight * sc * (0.6 if prec == "day-noyear" else 1.0) * (0.7 if p == "month" else 1.0)
        ctx = (before[-120:] + "【" + text[s:e] + "】" + after[:60]).replace("\n", " ¶ ")
        hint = extract_hint(text[max(0, s - 220):s], after, na)
        out.append(Candidate(kind, d, p, round(conf, 3), url, source_type, ctx, hint, translated))
    return out


# --------------------------------------------------------------------------- #
# Aggregation into intakes
# --------------------------------------------------------------------------- #
def _month_close(m1: int, m2: int, tol: int = 2) -> bool:
    diff = abs(m1 - m2) % 12
    return min(diff, 12 - diff) <= tol


def _resolve_hint_date(hint, ref: date) -> date:
    y, m = hint
    if y:
        return date(y, m, 1)
    y = ref.year
    while date(y, m, 1) < ref - timedelta(days=14):
        y += 1
    return date(y, m, 1)


def near_intake_words(context: str) -> bool:
    """Intake/admission wording right next to the date (context marks the date with 【】)."""
    before, _, rest = context.partition("【")
    after = rest.partition("】")[2]
    return bool(KW_INTAKE_WORDS.search(before[-45:] + " " + after[:25]))


def dedupe(cands: list[Candidate]) -> list[Candidate]:
    """Same sentence mirrored on several pages/snippets counts once (highest score kept)."""
    best: dict[tuple, Candidate] = {}
    for c in cands:
        key = (c.kind, c.date, re.sub(r"\W+", "", c.context.lower())[-90:])
        if key not in best or c.score > best[key].score:
            best[key] = c
    return list(best.values())


def build_intakes(cands: list[Candidate]) -> tuple[list[Intake], bool]:
    """Learn the intake pattern (start-month slots), then fill a slot x year grid.

    Returns every grid cell from two years back to two years ahead; cells with no
    evidence are kept (found=False) so the report can say "not found / not announced".
    """
    cands = dedupe(cands)
    starts = [c for c in cands if c.kind == "start"]
    # labelled deadlines ("Winter semester 2026/27 ... 15 July 2026") vote for a start month
    for c in cands:
        if c.kind in ("deadline", "open") and c.hint:
            vd = _resolve_hint_date(c.hint, c.date)
            if vd > c.date - timedelta(days=14) and (vd - c.date).days < 400:
                starts.append(Candidate("start", vd, "month", round(c.score * 0.35, 3), c.url,
                                        c.source_type, c.context, c.hint, c.translated, virtual=True))
    if not starts and not any(c.kind == "pattern" for c in cands):
        return [], False

    # 1. slots come mainly from intake evidence ("September 2026 intake", labelled deadlines);
    #    plain term starts ("Spring semester begins") count at reduced weight
    def slot_weight(c):
        return 1.0 if c.virtual or near_intake_words(c.context) else 0.5
    # each URL's say is capped so one big table can't define the pattern on its own
    per_url, per_url_intake = defaultdict(float), defaultdict(float)
    for c in starts:
        per_url[(c.date.month, c.url)] += c.score * slot_weight(c)
        if slot_weight(c) == 1.0:
            per_url_intake[(c.date.month, c.url)] += c.score
    for c in cands:
        if c.kind == "pattern":
            per_url[(c.date.month, c.url)] += c.score * 0.6
            per_url_intake[(c.date.month, c.url)] += c.score * 0.6

    def smoothed(pu):
        sup = defaultdict(float)
        for (month, _), v in pu.items():
            sup[month] += min(v, 1.5)
        return {m: sum(v for k, v in sup.items() if _month_close(k, m, 1)) for m in range(1, 13)}

    smooth, smooth_intake = smoothed(per_url), smoothed(per_url_intake)
    top, top_intake = max(smooth.values()), max(smooth_intake.values())
    slots: list[int] = []
    for m in sorted((m for m in smooth if smooth[m] > 0), key=lambda m: -smooth[m]):
        if any(_month_close(m, x, 2) for x in slots) or smooth[m] < 0.4 * top:
            continue
        # a secondary slot needs intake wording; plain term starts (a summer session) aren't enough
        if slots and top_intake > 0 and smooth_intake[m] < 0.35 * top_intake:
            continue
        slots.append(m)
    slots = slots[:3]
    log.debug("  slot support %s | intake-labelled %s -> %s",
              {m: round(v, 1) for m, v in smooth.items() if v},
              {m: round(v, 1) for m, v in smooth_intake.items() if v}, slots)
    # a slot's month = the month real (dated, non-inferred) start dates most often fall in
    real = defaultdict(float)
    for c in starts:
        if not c.virtual and c.precision == "day":
            real[c.date.month] += c.score
    adjusted: list[int] = []
    for m in slots:  # slots are in support order; drop any that collapse onto a stronger one
        m2 = max((k for k in real if _month_close(k, m, 1)), key=lambda k: real[k], default=m)
        if not any(_month_close(m2, x, 2) for x in adjusted):
            adjusted.append(m2)
    slots = sorted(adjusted)

    # 2. grid of (year, slot) cells, each anchored on the middle of its month
    cells = [Intake(start=date(y, m, 15), start_precision="nominal", start_score=0.0, slot=date(y, m, 15))
             for y in range(TODAY.year - 2, TODAY.year + 3) for m in slots]
    for c in starts:
        cell = min(cells, key=lambda it: abs((it.start - c.date).days))
        if abs((cell.start - c.date).days) <= 50:
            cell.cands.append(c)

    for it in cells:
        if not it.cands:
            continue
        by_date = defaultdict(float)
        for c in it.cands:
            by_date[(c.date, c.precision)] += c.score
        real_day = {k: v for k, v in by_date.items() if k[1] == "day" and
                    any(c.date == k[0] and c.precision == "day" and not c.virtual for c in it.cands)}
        (d, prec), _ = max((real_day or by_date).items(), key=lambda kv: (kv[1], -kv[0][0].toordinal()))
        it.start, it.start_precision = d, prec
        it.start_score = sum(by_date.values())
        it.inferred = all(c.virtual for c in it.cands)
        it.found = True

    # 3. attach deadlines / opening dates to the intake they belong to
    for c in cands:
        if c.kind not in ("deadline", "open"):
            continue
        target = None
        if c.hint:
            hd = _resolve_hint_date(c.hint, c.date)
            options = [it for it in cells if abs((it.start - hd).days) <= 62]
            target = min(options, key=lambda it: abs((it.start - hd).days), default=None)
            if target is None:
                continue  # labelled for an intake this university doesn't seem to have
        else:
            options = [it for it in cells if 0 <= (it.start - c.date).days <= 300]
            target = min(options, key=lambda it: (it.start - c.date).days, default=None)
        if target is not None and c.date <= target.start:
            (target.deadlines if c.kind == "deadline" else target.opens).append(c)
            target.found = True
    for it in cells:  # a "deadline" days before classes is usually course/enrolment admin
        for c in it.deadlines:
            if (it.start - c.date).days < 30:
                c.score = round(c.score * 0.35, 3)
    return cells, len(slots) >= 2


def best_date(cands: list[Candidate]):
    if not cands:
        return None, 0.0, []
    votes = defaultdict(float)
    for c in cands:
        votes[(c.date, c.precision)] += c.score
    top = max(votes.values())
    ranked = sorted(votes.items(), key=lambda kv: (-(kv[0][1] == "day" and kv[1] >= 0.5 * top),
                                                   -kv[1], kv[0][0]))
    (d, prec), score = ranked[0]
    others = [fmt_date(k[0], k[1]) for k, _ in ranked[1:4]]
    return (d, prec), score, others


def fmt_date(d: date | None, precision: str = "day") -> str:
    if d is None or precision == "nominal":
        return ""
    return d.strftime("%Y-%m-%d") if precision == "day" else d.strftime("%Y-%m") + " (month only)"


def classify(cells: list[Intake], multi: bool) -> list[tuple[str, Intake | None]]:
    """Pick Previous / Current / Next cells relative to today."""
    if not cells:
        labels = ["Previous-2", "Previous-1"] if multi else ["Previous"]
        return [(lbl, None) for lbl in labels + ["Current", "Next"]]
    cells = sorted(cells, key=lambda it: it.start)
    # current = the intake starting nearest to today (within ~4 months), else the upcoming one
    near = [i for i, it in enumerate(cells) if abs((it.start - TODAY).days) <= 120]
    if near:
        cur = min(near, key=lambda i: abs((cells[i].start - TODAY).days))
    else:
        cur = next((i for i, it in enumerate(cells) if it.start > TODAY), len(cells) - 1)
    n_prev = 2 if multi else 1
    labels = ["Previous-2", "Previous-1"] if multi else ["Previous"]
    prev = [cells[i] if i >= 0 else None for i in range(cur - n_prev, cur)]
    rows = list(zip(labels, prev))
    rows.append(("Current", cells[cur]))
    rows.append(("Next", cells[cur + 1] if cur + 1 < len(cells) else None))
    return rows


# --------------------------------------------------------------------------- #
# Pipeline per university
# --------------------------------------------------------------------------- #
class UniversityScraper:
    def __init__(self, country: str, uni: str, max_pages: int = 25, workers: int = 8,
                 aliases: list[str] | None = None, acronym: str = ""):
        self.country = country.strip()
        self.uni = uni.strip()
        self.aliases = list(dict.fromkeys([self.uni] + (aliases or [])))
        self.max_pages = max_pages
        self.workers = workers
        self.langs = local_languages(self.country, self.aliases)
        self.mdy = self.country.lower() in MDY_COUNTRIES
        self.na = self.mdy or self.country.lower() == "canada"
        self.acronym = acronym or "".join(w[0] for w in re.findall(r"[A-Za-z]+", self.uni)
                                          if w.lower() not in {"of", "the", "and", "for", "de", "at"}).upper()
        self.official: str | None = None
        self.sources: list[dict] = []
        self.cands: list[Candidate] = []
        self.seen_urls: set[str] = set()

    # -- helpers ------------------------------------------------------------
    def is_official(self, url: str) -> bool:
        dom = domain_of(url)
        return bool(self.official) and (dom == self.official or dom.endswith("." + self.official))

    def queries(self, second_pass: bool = False) -> list[str]:
        u, y = self.uni, TODAY.year
        if second_pass:
            q = [f"{u} {y - 1} intake application deadline", f"{u} academic calendar {y - 1}-{y}",
                 f"{u} {y - 1} semester start date", f"{u} admission dates {y - 1}"]
            if self.official:
                q.append(f"site:{self.official} {y - 1} application deadline")
        else:
            q = [f"{u} intake dates {y} {y + 1}", f"{u} application deadline {y}",
                 f"{u} application deadline {y + 1}", f"{u} semester start date {y}",
                 f"{u} academic calendar {y}-{y + 1}", f"{u} {self.country} intakes admission dates"]
            if self.official:
                q += [f"site:{self.official} application deadline",
                      f"site:{self.official} academic calendar",
                      f"site:{self.official} semester dates start"]
        phrases = ["application deadline", "semester start date", "academic calendar"]
        yy = y - 1 if second_pass else y
        for i, lang in enumerate(self.langs):
            local = [translate(p, target=lang, source="en") for p in phrases[: 3 if i == 0 else 1]]
            q += [f"{u} {p} {yy}" for p in local if p]
        return q

    def url_priority(self, url: str, title: str) -> int:
        blob = (url + " " + title).lower()
        p = 0
        if self.is_official(url):
            p += 3
        if re.search(r"admission|intake|deadline|calendar|dates|apply|semester|term|bewerb|fristen"
                     r"|termine|inscri|candidat|calend|admis|aanmeld|frist", blob):
            p += 2
        if any(a in url for a in AGGREGATORS):
            p += 1
        if url.lower().endswith(".pdf"):
            p += 1
        return p

    # -- stages -------------------------------------------------------------
    def gather_urls(self, queries: list[str]) -> list[tuple[str, str]]:
        results = []
        for q in queries:
            res = search(q)
            log.info("  search: %-70s -> %d results", q[:70], len(res))
            for r in res:
                href, body, title = r.get("href", ""), r.get("body", ""), r.get("title", "")
                if not href.startswith("http") or any(s in domain_of(href) for s in SKIP_DOMAINS):
                    continue
                # search-engine snippets are evidence too (useful when the page is blocked)
                body = SNIPPET_MID.sub(" ", SNIPPET_PREFIX.sub("", body))
                if body and (self.is_official(href) or mentions_university(title + " " + body, self.aliases, self.acronym)):
                    w = 0.8 if self.is_official(href) else 0.5
                    snip = extract_candidates(title + "\n" + body, href, "search-snippet",
                                              w, self.mdy, False, self.na)
                    if not self.is_official(href) and not mentions_university(title, self.aliases, self.acronym):
                        snip = [c for c in snip if c.kind != "pattern"]
                    self.cands += snip
                results.append((href, title))
            time.sleep(0.8)
        uniq = {}
        for href, title in results:
            if href not in self.seen_urls:
                uniq.setdefault(href.split("#")[0], title)
        ranked = sorted(uniq.items(), key=lambda kv: -self.url_priority(*kv))
        return ranked[: self.max_pages]

    def process_page(self, url: str, archive_ts: date | None = None) -> tuple[dict, list[Candidate]]:
        rec = {"University": self.uni, "URL": url, "Status": "", "Method": "live",
               "Language": "", "Translated": False, "Dates found": 0}
        official = self.is_official(url)
        if archive_ts is None:
            status, body, ctype, final = http_get(url)
        else:
            status, body, ctype, final = "skip", b"", "", url
        if status != "ok" and (official or status.startswith("blocked") or archive_ts):
            wb = wayback_get(url, archive_ts)
            if wb:
                body, ctype, final = wb
                status, rec["Method"] = "ok", "wayback-archive"
        rec["Status"] = status
        if status != "ok":
            return rec, []
        is_pdf = "pdf" in ctype or body[:5] == b"%PDF-" or url.lower().endswith(".pdf")
        if is_pdf:
            title, text = "", pdf_to_text(body)
            rec["Method"] += "+pdf"
        else:
            title, text = html_to_text(body)
        if len(text) < 200:
            rec["Status"] = "empty (JS-rendered or no text)"
            return rec, []
        lang = detect_lang(text)
        rec["Language"] = lang
        text = annotate_table_years(text)
        work_text, translated = text, False
        if lang != "en":
            work_text = translate_relevant(text)
            translated = rec["Translated"] = True
            title = translate(title) if title else title
        if not official and not mentions_university(title + "\n" + text + "\n" + work_text,
                                                    self.aliases, self.acronym):
            rec["Status"] = "skipped (page not about this university)"
            return rec, []
        if official:
            stype, w = ("official-archive", 1.8) if "wayback" in rec["Method"] else ("official", 2.0)
        else:
            stype, w = "third-party", 1.0
        path = urlparse(url).path.lower()
        if any(re.search(rf"\b{re.escape(k)}\b", path.replace("-", " ").replace("/", " "))
               for k in FOREIGN_CAMPUS_WORDS if k != self.country.lower()):
            w *= 0.3
        if ADMIN_PATH.search(path):  # census / exam / fee calendars list term dates, not intakes
            w *= 0.4
        cands = extract_candidates(work_text, final if "wayback" not in rec["Method"] else url,
                                   stype, w, self.mdy, translated, self.na)
        if not official and not mentions_university(title, self.aliases, self.acronym):
            # a generic "Intakes in France" article describes the country, not this university
            cands = [c for c in cands if c.kind != "pattern"]
        rec["Dates found"] = len(cands)
        return rec, cands

    def crawl(self, urls: list[str], archive_ts: date | None = None):
        with cf.ThreadPoolExecutor(self.workers) as ex:
            futs = {ex.submit(self.process_page, u, archive_ts): u for u in urls}
            for f in cf.as_completed(futs):
                try:
                    rec, cands = f.result()
                except Exception as e:  # noqa: BLE001
                    rec, cands = {"University": self.uni, "URL": futs[f], "Status": f"error {e}"}, []
                self.seen_urls.add(futs[f])
                self.sources.append(rec)
                self.cands += cands
                log.info("  %-28s %3d dates  %s", rec["Status"][:28], len(cands), futs[f][:90])

    def run(self) -> dict:
        log.info("=== %s (%s) ===", self.uni, self.country)
        self.official = find_official_domain(self.aliases, self.country)
        log.info("  official domain: %s | local language(s): %s", self.official, ", ".join(self.langs) or "en")

        urls = self.gather_urls(self.queries())
        self.crawl([u for u, _ in urls])
        cells, multi = build_intakes(self.cands)
        rows = classify(cells, multi)

        if not any(lbl.startswith("Previous") and it and it.found for lbl, it in rows):
            log.info("  no previous intake yet -> second pass (last year's dates + web archive)")
            urls2 = self.gather_urls(self.queries(second_pass=True))
            self.crawl([u for u, _ in urls2])
            # archived copies (≈1 year old) of official pages that contained dates
            dated_official = [s["URL"] for s in self.sources
                              if s.get("Dates found") and self.is_official(s["URL"])][:6]
            if dated_official:
                self.crawl(dated_official, archive_ts=TODAY - timedelta(days=330))
            cells, multi = build_intakes(self.cands)
            rows = classify(cells, multi)

        return {"rows": rows, "multi": multi}


# --------------------------------------------------------------------------- #
# Excel output
# --------------------------------------------------------------------------- #
def intake_name(it: Intake) -> str:
    return (it.slot or it.start).strftime("%B %Y") + " intake"


def to_rows(country: str, uni: str, official: str | None, result: dict, meta: dict | None = None):
    long_rows = []
    for label, it in result["rows"]:
        base = {"Country": country, "University": uni, **(meta or {}), "Official Website": official or "",
                "Intake Pattern": "Multiple intakes/year" if result["multi"] else "Single intake/year",
                "Intake Category": label}
        if it is None or not it.found:
            upcoming = it is not None and it.start > TODAY
            status = "Not announced yet" if upcoming or (it is None and label == "Next") else "Not found"
            long_rows.append({**base, "Intake": intake_name(it) if it else "", "Intake Start Date": "",
                              "Application Opens": "",
                              "Application Deadline": "", "Other Deadlines Seen": "",
                              "Confidence": "", "Status": status, "Sources": ""})
            continue
        dl, dl_score, dl_other = best_date(it.deadlines)
        op, _, _ = best_date(it.opens)
        total = it.score if it.start_precision != "nominal" else sum(c.score for c in it.deadlines)
        conf = "High" if total >= 3 else "Medium" if total >= 1.2 else "Low"
        notes = []
        if it.start_precision == "nominal":
            notes.append("start date not found")
        elif it.inferred:
            notes.append("start month inferred from intake label")
        if not dl:
            notes.append("deadline not found")
        all_c = sorted(it.cands + it.deadlines + it.opens, key=lambda c: -c.score)
        srcs = list(dict.fromkeys(c.url for c in all_c if c.source_type.startswith("official")))
        srcs += [u for u in dict.fromkeys(c.url for c in all_c) if u not in srcs]
        long_rows.append({
            **base,
            "Intake": intake_name(it),
            "Intake Start Date": fmt_date(it.start, it.start_precision),
            "Application Opens": fmt_date(*op) if op else "",
            "Application Deadline": fmt_date(*dl) if dl else "",
            "Other Deadlines Seen": ", ".join(dl_other),
            "Confidence": conf,
            "Status": "Found" + (f" ({'; '.join(notes)})" if notes else ""),
            "Sources": "\n".join(srcs[:5]),
        })
    return long_rows


def write_excel(path: str, long_rows: list[dict], evidence: list[dict], sources: list[dict]):
    df = pd.DataFrame(long_rows)
    summary = []
    for (country, uni), g in df.groupby(["Country", "University"], sort=False):
        row = {"Country": country, "University": uni,
               **({"Institution Type": g["Institution Type"].iloc[0]} if "Institution Type" in g else {}),
               "Official Website": g["Official Website"].iloc[0],
               "Intake Pattern": g["Intake Pattern"].iloc[0]}
        for _, r in g.iterrows():
            cat = r["Intake Category"]
            if r["Intake"]:
                row[f"{cat} Intake"] = r["Intake"]
                row[f"{cat} Start"] = r["Intake Start Date"]
                row[f"{cat} Deadline"] = r["Application Deadline"] or "not found"
            else:
                row[f"{cat} Intake"] = r["Status"]
        row["Previous intake captured?"] = "YES" if any(
            str(c).startswith("Previous") and i for c, i in zip(g["Intake Category"], g["Intake"])) else "NO"
        summary.append(row)

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(summary).to_excel(xw, sheet_name="Summary", index=False)
        df.to_excel(xw, sheet_name="Intakes", index=False)
        pd.DataFrame(evidence).to_excel(xw, sheet_name="Evidence", index=False)
        pd.DataFrame(sources).to_excel(xw, sheet_name="Sources", index=False)
        from openpyxl.styles import Alignment, Font, PatternFill
        head_fill = PatternFill("solid", fgColor="1F4E78")
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for col in ws.columns:
                width = max((len(str(c.value)) if c.value is not None else 0) for c in col[:300])
                ws.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), 70)
                for c in col[1:]:
                    c.alignment = Alignment(wrap_text=True, vertical="top")


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
MINUTES_PER_UNI = 3


def job(country: str, name: str, **extra) -> dict:
    return {"country": country, "name": name, "aliases": [name], "acronym": "", "type": "",
            "raw_name": name, **extra}


def ask(prompt: str) -> str:
    try:
        return input(prompt)
    except EOFError:
        return ""


def pick_from_registry(registry: dict, country: str | None, only: str | None) -> list[dict]:
    """Country menu -> university list -> subset selection. Returns scrape jobs."""
    countries = list(registry)
    if country:
        chosen = match_country(country, countries)
        if not chosen:
            sys.exit(f"'{country}' is not a tab in the registry. Tabs: {', '.join(countries)}")
    else:
        print(f"\nCountries in registry ({len(countries)}):")
        width = max(len(c) for c in countries) + 6
        cells = [f"{i:>2}. {c} ({sum(1 for e in registry[c] if not e.duplicate_of)})"
                 for i, c in enumerate(countries, 1)]
        for i in range(0, len(cells), 3):
            print("  " + "".join(cell.ljust(width + 4) for cell in cells[i:i + 3]))
        chosen = None
        while not chosen:
            chosen = match_country(ask("\nSelect country (number or name): "), countries)
            if not chosen:
                print("  not recognised, try again")

    entries = registry[chosen]
    dups = [e for e in entries if e.duplicate_of]
    unis = [e for e in entries if not e.duplicate_of]
    print(f"\n{chosen}: {len(unis)} universities"
          + (f" ({len(dups)} duplicate row(s) skipped)" if dups else ""))
    for d in dups:
        print(f"   skipped duplicate: '{d.raw_name}' (same as '{d.duplicate_of}')")
    for i, e in enumerate(unis, 1):
        extra = f"  [{e.acronym}]" if e.acronym else ""
        alt = f"  aka {', '.join(e.aliases[1:])}" if len(e.aliases) > 1 else ""
        print(f"  {i:>3}. {e.name}{extra}{alt}  — {e.type}")
    if only is None and sys.stdin.isatty():
        only = ask("\nWhich universities? [Enter = all, or e.g. 1-5,8,12]: ")
    picked = [unis[i] for i in parse_selection(only or "", len(unis))]
    return [job(chosen, e.name, aliases=e.aliases, acronym=e.acronym, type=e.type, raw_name=e.raw_name)
            for e in picked]


def load_jobs(args, registry: dict | None) -> list[dict]:
    if args.input:
        df = pd.read_excel(args.input) if args.input.endswith((".xlsx", ".xls")) else pd.read_csv(args.input)
        df.columns = [c.strip().lower() for c in df.columns]
        return [job(str(r["country"]), str(r["university"])) for _, r in df.iterrows()]
    if args.country and args.university:
        return [job(args.country, u) for u in args.university]
    if registry:
        return pick_from_registry(registry, args.country, args.only)
    jobs = []
    print("Enter country and university (blank university to finish).")
    while True:
        country = ask("Country: ").strip() or (jobs[-1]["country"] if jobs else "")
        uni = ask("University: ").strip()
        if not uni:
            break
        jobs.append(job(country, uni))
    return jobs


def load_previous(path: str) -> tuple[list[dict], list[dict], list[dict]]:
    """Rows already saved in an earlier (possibly interrupted) run of the same output file."""
    try:
        sheets = pd.read_excel(path, sheet_name=["Intakes", "Evidence", "Sources"])
    except Exception:  # noqa: BLE001
        return [], [], []
    return tuple(sheets[k].fillna("").to_dict("records") for k in ("Intakes", "Evidence", "Sources"))


def scrape_one(j: dict, max_pages: int) -> tuple[list[dict], list[dict], list[dict], bool]:
    s = UniversityScraper(j["country"], j["name"], max_pages=max_pages,
                          aliases=j["aliases"], acronym=j["acronym"])
    try:
        res = s.run()
    except Exception as e:  # noqa: BLE001
        log.exception("failed on %s: %s", j["name"], e)
        res = {"rows": [("Previous", None), ("Current", None), ("Next", None)], "multi": False}
    meta = {"Name in Registry": j["raw_name"], "Institution Type": j["type"]}
    rows = to_rows(j["country"], j["name"], s.official, res, meta)
    evidence = [{"Country": j["country"], "University": j["name"], "Type": c.kind,
                 "Date": (c.date.strftime("%B") + " (intake pattern)") if c.kind == "pattern"
                 else fmt_date(c.date, c.precision), "Intake hint": str(c.hint or ""),
                 "Score": c.score, "Source type": c.source_type,
                 "Translated": c.translated, "URL": c.url, "Context": c.context}
                for c in sorted(s.cands, key=lambda c: (c.kind, c.date))]
    for r in rows:
        log.info("  -> %s | %-10s %-22s start=%-12s deadline=%s", j["name"][:30], r["Intake Category"],
                 r["Intake"], r["Intake Start Date"], r["Application Deadline"] or r["Status"])
    return rows, evidence, s.sources, any(str(r["Intake Category"]).startswith("Previous") and
                                          str(r["Status"]).startswith("Found") for r in rows)


def main():
    ap = argparse.ArgumentParser(
        description="Scrape university intake dates into Excel. With no arguments it lists the "
                    "countries in the registry workbook and asks which one to scrape.")
    ap.add_argument("-c", "--country", help="country tab to scrape (all its universities unless -u is given)")
    ap.add_argument("-u", "--university", action="append", help="single university (repeatable)")
    ap.add_argument("-i", "--input", help="CSV/XLSX with columns: country, university")
    ap.add_argument("-r", "--registry", default=str(DEFAULT_REGISTRY),
                    help="workbook with one tab per country (default: %(default)s)")
    ap.add_argument("--only", help="subset of the country's list, e.g. '1-10,15' (skips the prompt)")
    ap.add_argument("-o", "--output", help="output workbook (default: intake_dates_<Country>.xlsx)")
    ap.add_argument("--fresh", action="store_true", help="ignore universities already in the output file")
    ap.add_argument("-p", "--parallel", type=int, default=1, help="universities scraped at the same time")
    ap.add_argument("--max-pages", type=int, default=25, help="pages fetched per university per pass")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args()
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(message)s", stream=sys.stdout)
    for noisy in ("urllib3", "charset_normalizer", "primp", "httpx", "ddgs"):
        logging.getLogger(noisy).setLevel(logging.WARNING)

    registry = None
    if not args.input and not (args.country and args.university):
        if Path(args.registry).exists():
            registry = load_registry(args.registry)
        elif args.country or args.registry != str(DEFAULT_REGISTRY):
            sys.exit(f"registry workbook not found: {args.registry}")

    jobs = load_jobs(args, registry)
    if not jobs:
        ap.error("no universities selected")
    countries = sorted({j["country"] for j in jobs})
    output = args.output or (f"intake_dates_{countries[0].replace(' ', '_')}.xlsx"
                             if len(countries) == 1 else "intake_dates.xlsx")

    long_rows, evidence, sources = ([], [], []) if args.fresh else load_previous(output)
    done = {(str(r.get("Country")), str(r.get("University"))) for r in long_rows}
    todo = [j for j in jobs if (j["country"], j["name"]) not in done]
    if len(todo) < len(jobs):
        print(f"\nResuming {output}: {len(jobs) - len(todo)} already done, {len(todo)} to go "
              f"(use --fresh to redo them)")
    if not todo:
        print("Nothing left to do.")
        return
    eta = len(todo) * MINUTES_PER_UNI / max(1, args.parallel)
    print(f"\nScraping {len(todo)} universit{'y' if len(todo) == 1 else 'ies'} "
          f"(~{eta:.0f} min{', ' + str(args.parallel) + ' in parallel' if args.parallel > 1 else ''}). "
          f"Saving to {output} after each one. Ctrl+C stops safely; run again to resume.\n")

    lock = threading.Lock()
    captured = 0
    unsaved = False

    def save() -> bool:
        try:
            write_excel(output, long_rows, evidence, sources)
            return True
        except PermissionError:  # usually: the workbook is open in Excel (Windows locks it)
            log.warning("!! cannot write %s (is it open in Excel?) - will retry after the next university",
                        output)
            return False

    def collect(result):
        nonlocal captured, unsaved
        rows, ev, src, has_prev = result
        with lock:
            long_rows.extend(rows)
            evidence.extend(ev)
            sources.extend(src)
            captured += has_prev
            unsaved = not save()

    finished = 0
    try:
        with cf.ThreadPoolExecutor(max(1, args.parallel)) as ex:
            futs = {ex.submit(scrape_one, j, args.max_pages): j for j in todo}
            for f in cf.as_completed(futs):
                collect(f.result())
                finished += 1
                log.info("[%d/%d done] %s saved", finished, len(todo), futs[f]["name"])
    except KeyboardInterrupt:
        print(f"\nStopped. {finished} finished universit{'y is' if finished == 1 else 'ies are'} saved "
              f"in {output}; run the same command again to continue.")
        os._exit(1)  # don't wait for in-flight page downloads

    for _ in range(6):
        if not unsaved:
            break
        print(f"Close {output} in Excel - retrying in 10 s ...")
        time.sleep(10)
        unsaved = not save()
    if unsaved:
        output = output.replace(".xlsx", f"_backup_{int(time.time())}.xlsx")
        write_excel(output, long_rows, evidence, sources)
    print(f"\nSaved {output}: previous intake captured for {captured}/{len(todo)} universities this run.")


if __name__ == "__main__":
    main()
